# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
"""Versioned, evaluator-only criteria and non-executing artifact checks.

Rules are frozen before grading. Spreadsheet formulas are never executed: a
missing cached result is explicitly unconfirmed (or fails the declared missing
policy), not silently replaced by a generated value.
"""

from __future__ import annotations

import fnmatch
import json
import zipfile
from decimal import ROUND_CEILING, ROUND_FLOOR, ROUND_HALF_EVEN, ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Mapping
from xml.etree.ElementTree import ParseError

import openpyxl
import yaml
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from .artifacts import sha256_bytes
from .errors import ConfigError

_ROUNDING = {"half_even": ROUND_HALF_EVEN, "half_up": ROUND_HALF_UP, "ceiling": ROUND_CEILING, "floor": ROUND_FLOOR}
_KINDS = {"file_exists", "text_contains", "json_number", "xlsx_cell", "xlsx_sum", "sheet_exists"}


def _decimal(value: Any) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise ValueError("expected a finite decimal number")
    result = Decimal(str(value))
    if not result.is_finite():
        raise ValueError("expected a finite decimal number")
    return result


def load_criteria(value: Any, base: Path) -> dict[str, Any] | None:
    """Resolve a YAML/JSON file or an inline policy; validate without grading."""
    if value is None:
        return None
    if isinstance(value, str):
        source = Path(value).expanduser()
        source = source if source.is_absolute() else base / source
        try:
            value = yaml.safe_load(source.read_bytes())
        except (OSError, yaml.YAMLError) as exc:
            raise ConfigError(f"cannot read grading criteria: {source}: {exc}") from exc
    if not isinstance(value, dict) or set(value) != {"version", "benchmark", "policy", "tasks"}:
        raise ConfigError("criteria requires exactly version, benchmark, policy and tasks")
    if not isinstance(value["version"], str) or not value["version"].strip():
        raise ConfigError("criteria.version must be a non-empty version identifier")
    if value["benchmark"] not in {"gdpval", "gsm8k"}:
        raise ConfigError("criteria.benchmark must be gdpval or gsm8k")
    policy = value["policy"]
    policy_keys = {
        "arithmetic",
        "rounding",
        "decimal_places",
        "absolute_tolerance",
        "relative_tolerance",
        "missing",
        "units",
    }
    if not isinstance(policy, dict) or set(policy) != policy_keys:
        raise ConfigError(
            "criteria.policy must declare arithmetic, rounding, decimal_places, absolute_tolerance, relative_tolerance, missing and units"
        )
    if policy["arithmetic"] != "decimal" or policy["rounding"] not in {*_ROUNDING, "none"}:
        raise ConfigError("criteria.policy requires decimal arithmetic and a supported rounding mode")
    places = policy["decimal_places"]
    if isinstance(places, bool) or not isinstance(places, int) or not 0 <= places <= 12:
        raise ConfigError("criteria.policy.decimal_places must be an integer from 0 to 12")
    try:
        if any(_decimal(policy[k]) < 0 for k in ("absolute_tolerance", "relative_tolerance")):
            raise ValueError("negative tolerance")
    except (ValueError, InvalidOperation) as exc:
        raise ConfigError("criteria tolerances must be finite nonnegative numbers") from exc
    if policy["missing"] not in {"fail", "unconfirmed"}:
        raise ConfigError("criteria.policy.missing must be fail or unconfirmed")
    if not isinstance(policy["units"], str) or not policy["units"].strip():
        raise ConfigError("criteria.policy.units must state the unit convention")
    tasks = value["tasks"]
    if not isinstance(tasks, dict) or not tasks:
        raise ConfigError("criteria.tasks must contain task IDs (or '*' for common rules)")
    for task_id, task in tasks.items():
        if (
            not isinstance(task_id, str)
            or not task_id
            or not isinstance(task, dict)
            or set(task) != {"mechanical", "ai"}
        ):
            raise ConfigError("each criteria task requires mechanical and ai lists")
        if not task["mechanical"] and not task["ai"]:
            raise ConfigError("each criteria task must declare at least one grading item")
        identifiers: set[str] = set()
        for category in ("mechanical", "ai"):
            rules = task[category]
            if not isinstance(rules, list):
                raise ConfigError(f"criteria.tasks.{task_id}.{category} must be a list")
            for rule in rules:
                if not isinstance(rule, dict) or not isinstance(rule.get("id"), str) or not rule["id"].strip():
                    raise ConfigError("every grading item requires an id")
                if rule["id"] in identifiers:
                    raise ConfigError("grading item IDs must be unique within a task")
                identifiers.add(rule["id"])
                if not isinstance(rule.get("description"), str) or not rule["description"].strip():
                    raise ConfigError("every grading item requires a description")
                if category == "ai":
                    if set(rule) != {"id", "description"}:
                        raise ConfigError("AI criteria accept only id and description; shared policy applies")
                    continue
                allowed = {
                    "id",
                    "description",
                    "kind",
                    "path",
                    "expected",
                    "sheet",
                    "cell",
                    "range",
                    "key",
                    "unit",
                    "unit_cell",
                }
                if set(rule) - allowed or rule.get("kind") not in _KINDS:
                    raise ConfigError("mechanical criterion has an unsupported kind or field")
                path = rule.get("path")
                if not isinstance(path, str) or not path or Path(path).is_absolute() or ".." in Path(path).parts:
                    raise ConfigError("mechanical paths must stay within the saved submission")
                kind = rule["kind"]
                required = {
                    "text_contains": ["expected"],
                    "json_number": ["key", "expected"],
                    "xlsx_cell": ["sheet", "cell", "expected"],
                    "xlsx_sum": ["sheet", "range", "expected"],
                    "sheet_exists": ["sheet"],
                }.get(kind, [])
                if any(k not in rule for k in required):
                    raise ConfigError(f"mechanical {kind} requires {', '.join(required)}")
                if kind in {"json_number", "xlsx_sum"}:
                    try:
                        _decimal(rule["expected"])
                    except (InvalidOperation, ValueError) as exc:
                        raise ConfigError(f"{kind}.expected must be numeric") from exc
                if ("unit" in rule) != ("unit_cell" in rule) or (
                    "unit" in rule and kind not in {"xlsx_cell", "xlsx_sum"}
                ):
                    raise ConfigError("spreadsheet unit checks require both unit and unit_cell")
    # A detached JSON value prevents caller mutation and rejects unsafe/non-JSON YAML values.
    try:
        return json.loads(json.dumps(value, allow_nan=False))  # type: ignore[no-any-return]
    except (ValueError, TypeError) as exc:
        raise ConfigError("grading criteria must contain finite JSON-compatible values") from exc


def criteria_hash(criteria: Mapping[str, Any] | None) -> str | None:
    if criteria is None:
        return None
    return sha256_bytes(json.dumps(criteria, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode())


def criteria_for_task(criteria: Mapping[str, Any] | None, task_id: str, benchmark: str) -> dict[str, Any] | None:
    if criteria is None:
        return None
    if criteria["benchmark"] != benchmark:
        raise ConfigError("grading criteria benchmark does not match the source run")
    tasks = criteria["tasks"]
    if task_id not in tasks and "*" not in tasks:
        raise ConfigError(f"grading criteria do not cover task {task_id}")
    selected = [tasks[key] for key in ("*", task_id) if key in tasks]
    result = {"version": criteria["version"], "policy": criteria["policy"], "mechanical": [], "ai": []}
    for category in ("mechanical", "ai"):
        result[category] = [rule for task in selected for rule in task[category]]
    ids = [r["id"] for category in ("mechanical", "ai") for r in result[category]]
    if len(set(ids)) != len(ids):
        raise ConfigError(f"common and task-specific grading items collide for {task_id}")
    return result


def _numeric_check(observed: Any, expected: Any, policy: Mapping[str, Any]) -> tuple[bool, dict[str, Any]]:
    actual, target = _decimal(observed), _decimal(expected)
    if policy["rounding"] != "none":
        quantizer = Decimal(1).scaleb(-policy["decimal_places"])
        mode = _ROUNDING[policy["rounding"]]
        actual, target = actual.quantize(quantizer, rounding=mode), target.quantize(quantizer, rounding=mode)
    tolerance = max(_decimal(policy["absolute_tolerance"]), abs(target) * _decimal(policy["relative_tolerance"]))
    return abs(actual - target) <= tolerance, {
        "observed": str(actual),
        "expected": str(target),
        "allowed_error": str(tolerance),
    }


def mechanical_findings(submission: Path, task_criteria: Mapping[str, Any] | None) -> list[dict[str, Any]]:
    """Inspect a captured submission without executing formulas, scripts or macros."""
    if task_criteria is None:
        return []
    root = submission.resolve()
    policy = task_criteria["policy"]
    findings: list[dict[str, Any]] = []
    for rule in task_criteria["mechanical"]:
        finding: dict[str, Any] = {
            "criterion_id": rule["id"],
            "description": rule["description"],
            "method": "mechanical",
            "status": "unconfirmed",
            "evidence": {"path": rule["path"]},
        }
        try:
            paths = sorted(
                p
                for p in root.rglob("*")
                if p.is_file() and fnmatch.fnmatch(p.relative_to(root).as_posix(), rule["path"])
            )
            if any(not p.resolve().is_relative_to(root) or p.is_symlink() for p in paths):
                raise ValueError("submission path escapes the captured artifact root")
            if not paths:
                raise FileNotFoundError("required file is missing")
            if rule["kind"] == "file_exists":
                finding.update(status="pass", observed=[p.relative_to(root).as_posix() for p in paths])
            elif len(paths) != 1:
                raise ValueError("selector matches multiple files; result is ambiguous")
            elif rule["kind"] == "text_contains":
                if paths[0].stat().st_size > 1024 * 1024:
                    raise ValueError("text artifact exceeds the 1 MiB inspection limit")
                finding.update(
                    status="pass" if str(rule["expected"]) in paths[0].read_text() else "fail",
                    expected=rule["expected"],
                )
            else:
                if rule["kind"] == "json_number":
                    if paths[0].stat().st_size > 1024 * 1024:
                        raise ValueError("JSON artifact exceeds the 1 MiB inspection limit")
                    observed = json.loads(paths[0].read_text())
                    for key in str(rule["key"]).split("."):
                        observed = observed[key]
                else:
                    with zipfile.ZipFile(paths[0]) as archive:
                        if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
                            raise ValueError("spreadsheet exceeds the 64 MiB expanded inspection limit")
                    book = openpyxl.load_workbook(paths[0], data_only=True, read_only=True, keep_links=False)
                    try:
                        if rule["sheet"] not in book.sheetnames:
                            raise KeyError("required sheet is missing")
                        sheet = book[rule["sheet"]]
                        finding["evidence"]["sheet"] = rule["sheet"]
                        if rule["kind"] == "sheet_exists":
                            finding["status"] = "pass"
                            findings.append(finding)
                            continue
                        if "unit" in rule:
                            unit = sheet[rule["unit_cell"]].value
                            finding["unit"] = {
                                "expected": rule["unit"],
                                "observed": str(unit),
                                "cell": rule["unit_cell"],
                            }
                            if unit is None:
                                raise KeyError("unit is missing")
                            if str(unit).strip() != str(rule["unit"]).strip():
                                finding.update(status="fail", reason="unit does not match the declared convention")
                                findings.append(finding)
                                continue
                        if rule["kind"] == "xlsx_cell":
                            observed = sheet[rule["cell"]].value
                            finding["evidence"]["cell"] = rule["cell"]
                        else:
                            a, b, c, d = range_boundaries(rule["range"])
                            if a is None or b is None or c is None or d is None:
                                raise ValueError("sum range must specify bounded cells")
                            if (c - a + 1) * (d - b + 1) > 50000:
                                raise ValueError("sum range exceeds 50000 cells")
                            values = [
                                cell.value
                                for row in sheet.iter_rows(min_row=b, max_row=d, min_col=a, max_col=c)
                                for cell in row
                            ]
                            if any(v is None for v in values):
                                raise KeyError("sum range contains missing or uncached formula values")
                            observed = sum((_decimal(v) for v in values), Decimal(0))
                            finding["evidence"]["range"] = rule["range"]
                    finally:
                        book.close()
                if observed is None:
                    raise KeyError("value is missing or the formula has no cached result")
                try:
                    passed, details = _numeric_check(observed, rule["expected"], policy)
                except (ValueError, InvalidOperation):
                    if rule["kind"] != "xlsx_cell" or not isinstance(rule["expected"], str):
                        raise
                    passed, details = (
                        observed == rule["expected"],
                        {"observed": str(observed), "expected": rule["expected"]},
                    )
                finding.update(status="pass" if passed else "fail", **details)
        except (FileNotFoundError, KeyError) as exc:
            finding.update(status=policy["missing"], reason=str(exc))
        except (
            OSError,
            ValueError,
            TypeError,
            InvalidOperation,
            IndexError,
            zipfile.BadZipFile,
            ParseError,
            InvalidFileException,
        ) as exc:
            finding.update(status="unconfirmed", reason=str(exc))
        findings.append(finding)
    return findings
