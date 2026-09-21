# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0

"""Derivative provenance rejects unbound bytes and cross-submission aliases."""

from __future__ import annotations

import base64
import copy
import json
import shutil
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from eval_harness.artifacts import file_manifest, manifest_hash, sha256_file, write_json, write_jsonl
from eval_harness.errors import ArtifactError
from eval_harness.inspection_lineage import find_derived_scopes, validate_derivations
from eval_harness.inspection_protocol import _check_ref, _shell_record_scopes, _verified_calls

_FRAME = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4z8AAAAMBAQDJ/pLvAAAAAElFTkSuQmCC"
)
_SOURCE = "submission_A/clip.mp4"
_DERIVATIVE = ".harness_evidence/scratch/frames/frame-0001.png"
_STREAM = ".harness_evidence/outputs/independent-process-uuid.stdout.log"
_CALL = "000001-journal-call-id"


def _put(workspace: Path, path: str, data: bytes) -> str:
    target = workspace / path
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)
    return sha256_file(target)


def _fixture(workspace: Path) -> dict[str, Any]:
    original = _put(workspace, _SOURCE, b"fixed synthetic original clip bytes")
    frame = _put(workspace, _DERIVATIVE, _FRAME)
    return {
        "schema_version": 1,
        "source_artifacts": [{"path": _SOURCE, "sha256": original}],
        "checks": [{"action": "Extract frame at 00:00:01.000", "observation": {"frame_count": 1}}],
        "derived_artifacts": [
            {"path": _DERIVATIVE, "sha256": frame, "source_paths": [_SOURCE], "location": "00:00:01.000"}
        ],
    }


def _calls(workspace: Path, record: dict[str, Any], *, stream: str = "stdout") -> dict[str, dict[str, Any]]:
    path = _STREAM.replace("stdout", stream)
    digest = _put(workspace, path, json.dumps(record).encode())
    return {
        _CALL: {
            "call_id": _CALL,
            "tool": "shell",
            "arguments": {"command": "known-fixture frame inspection"},
            "failure": None,
            "result": {
                "status": "completed",
                "returncode": 0,
                f"captured_{stream}_path": path,
                f"{stream}_sha256": digest,
            },
        }
    }


def _guard(path: Path, workspace: Path) -> set[str]:
    return _shell_record_scopes(path, workspace, method="content")


def _find(workspace: Path, record: dict[str, Any], calls: dict[str, dict[str, Any]]) -> set[str]:
    return find_derived_scopes(record["derived_artifacts"][0]["sha256"], workspace, calls, validate_record=_guard)


def test_known_frame_chain_uses_verified_journal_and_performed_check_guard(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    calls = _calls(tmp_path, record)
    receipt = calls[_CALL]
    root = tmp_path / ".harness_evidence"
    write_json(root / f"receipts/{_CALL}.json", receipt)
    common = {"call_id": _CALL, "tool": "shell", "arguments": receipt["arguments"]}
    write_jsonl(
        root / "events.jsonl",
        [
            {**common, "event": "started"},
            {
                **common,
                "event": "finished",
                "status": "completed",
                "failure": None,
                "output_sha256": sha256_file(root / f"receipts/{_CALL}.json"),
            },
        ],
    )
    entries = file_manifest(root)
    write_json(root / "manifest.json", {"files": entries})
    trusted = _verified_calls(tmp_path, manifest_hash(entries))
    assert _find(tmp_path, record, trusted) == {"submission_A"}
    derived = validate_derivations(record, tmp_path)[0]
    assert derived.path == _DERIVATIVE
    assert derived.source_paths == (_SOURCE,)
    assert derived.location == "00:00:01.000"
    # Valid hashes establish the recorded chain, not the synthetic extraction's truth.
    assert derived.scopes == frozenset({"submission_A"})


def test_stderr_record_and_distinct_shell_process_id_are_supported(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    assert _find(tmp_path, record, _calls(tmp_path, record, stream="stderr")) == {"submission_A"}


def test_legacy_record_and_absent_calls_confer_no_lineage(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    original = copy.deepcopy(record)
    del record["derived_artifacts"]
    assert validate_derivations(record, tmp_path) == ()
    assert _find(tmp_path, original, _calls(tmp_path, record)) == set()
    assert _find(tmp_path, original, {}) == set()


def test_unrelated_view_hash_is_not_accepted(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    unrelated = _put(tmp_path, ".harness_evidence/scratch/other.png", b"other frame bytes")
    assert find_derived_scopes(unrelated, tmp_path, _calls(tmp_path, record), validate_record=_guard) == set()


@pytest.mark.parametrize("target", ["original", "derivative", "stream"])
def test_tampered_bytes_are_rejected(tmp_path: Path, target: str) -> None:
    record = _fixture(tmp_path)
    calls = _calls(tmp_path, record)
    path = {"original": _SOURCE, "derivative": _DERIVATIVE, "stream": _STREAM}[target]
    (tmp_path / path).write_bytes(b"changed after the recorded hashes")
    with pytest.raises(ArtifactError, match="hash"):
        _find(tmp_path, record, calls)


@pytest.mark.parametrize("target", [_SOURCE, _DERIVATIVE, _STREAM])
def test_missing_captured_or_original_file_is_rejected(tmp_path: Path, target: str) -> None:
    record = _fixture(tmp_path)
    calls = _calls(tmp_path, record)
    (tmp_path / target).unlink()
    with pytest.raises(ArtifactError, match="missing"):
        _find(tmp_path, record, calls)


def test_wrong_original_hash_is_rejected(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    record["source_artifacts"][0]["sha256"] = "0" * 64
    with pytest.raises(ArtifactError, match="hash"):
        _find(tmp_path, record, _calls(tmp_path, record))


@pytest.mark.parametrize(
    "updates",
    [
        {"returncode": 1},
        {"returncode": False},
        {"status": "timeout"},
        {"status": "failed"},
    ],
)
def test_unsuccessful_shell_cannot_establish_derivative_scope(tmp_path: Path, updates: dict[str, Any]) -> None:
    record = _fixture(tmp_path)
    calls = _calls(tmp_path, record)
    calls[_CALL]["result"].update(updates)
    assert _find(tmp_path, record, calls) == set()


def test_failed_receipt_and_other_tool_cannot_establish_scope(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    calls = _calls(tmp_path, record)
    calls[_CALL]["failure"] = {"message": "process failed"}
    assert _find(tmp_path, record, calls) == set()
    calls[_CALL]["failure"] = None
    calls[_CALL]["tool"] = "inspect_document"
    assert _find(tmp_path, record, calls) == set()


def test_callback_failure_cannot_be_bypassed_by_valid_lineage(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    record["checks"] = []
    with pytest.raises(ArtifactError, match="checks"):
        _find(tmp_path, record, _calls(tmp_path, record))


def test_pairwise_record_only_credits_explicit_derivative_sources(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    # Equal original hashes still cannot transfer membership from A to B.
    digest_b = _put(tmp_path, "submission_B/clip.mp4", (tmp_path / _SOURCE).read_bytes())
    record["source_artifacts"].append({"path": "submission_B/clip.mp4", "sha256": digest_b})
    calls = _calls(tmp_path, record)
    assert _guard(tmp_path / _STREAM, tmp_path) == {"submission_A", "submission_B"}
    assert _find(tmp_path, record, calls) == {"submission_A"}
    record["derived_artifacts"][0]["source_paths"] = ["submission_B/clip.mp4"]
    assert _find(tmp_path, record, _calls(tmp_path, record)) == {"submission_B"}


def test_pairwise_different_frames_do_not_union_scopes(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    source_b = "submission_B/clip.mp4"
    path_b = ".harness_evidence/scratch/frames/b.png"
    record["source_artifacts"].append({"path": source_b, "sha256": _put(tmp_path, source_b, b"original B")})
    digest_b = _put(tmp_path, path_b, b"different B frame")
    record["derived_artifacts"].append(
        {"path": path_b, "sha256": digest_b, "source_paths": [source_b], "location": "00:00:02.000"}
    )
    calls = _calls(tmp_path, record)
    assert _find(tmp_path, record, calls) == {"submission_A"}
    assert find_derived_scopes(digest_b, tmp_path, calls, validate_record=_guard) == {"submission_B"}
    record["derived_artifacts"][0]["sha256"] = digest_b
    with pytest.raises(ArtifactError, match="hash"):
        _find(tmp_path, record, _calls(tmp_path, record))


@pytest.mark.parametrize(
    "paths",
    [["clip.mp4"], ["submission_B/clip.mp4"], [_SOURCE, _SOURCE], [], [123]],
)
def test_source_membership_requires_exact_declared_paths(tmp_path: Path, paths: list[Any]) -> None:
    record = _fixture(tmp_path)
    record["derived_artifacts"][0]["source_paths"] = paths
    with pytest.raises(ArtifactError, match="distinct originals"):
        validate_derivations(record, tmp_path)


@pytest.mark.parametrize(
    "path",
    [
        "frame-0001.png",
        "scratch/frame-0001.png",
        ".harness_evidence//scratch/frames/frame-0001.png",
        ".harness_evidence/scratch/frames/../frames/frame-0001.png",
        "./" + _DERIVATIVE,
    ],
)
def test_noncanonical_derivative_alias_is_rejected(tmp_path: Path, path: str) -> None:
    record = _fixture(tmp_path)
    if ".." not in path and "//" not in path and not path.startswith("./"):
        _put(tmp_path, path, _FRAME)
    record["derived_artifacts"][0]["path"] = path
    with pytest.raises(ArtifactError, match="path"):
        validate_derivations(record, tmp_path)


def test_absolute_and_embedded_submission_source_aliases_do_not_confer_scope(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    absolute = str(tmp_path / _SOURCE)
    record["source_artifacts"][0]["path"] = absolute
    record["derived_artifacts"][0]["source_paths"] = [absolute]
    with pytest.raises(ArtifactError, match="relative"):
        validate_derivations(record, tmp_path)
    alias = "reference_files/submission_A/clip.mp4"
    _put(tmp_path, alias, (tmp_path / _SOURCE).read_bytes())
    record["source_artifacts"][0]["path"] = alias
    record["derived_artifacts"][0]["source_paths"] = [alias]
    assert validate_derivations(record, tmp_path)[0].scopes == frozenset()


@pytest.mark.parametrize("target", [_SOURCE, _DERIVATIVE, _STREAM])
def test_symlink_aliases_are_rejected(tmp_path: Path, target: str) -> None:
    record = _fixture(tmp_path)
    calls = _calls(tmp_path, record)
    path = tmp_path / target
    actual = path.with_suffix(".actual")
    path.rename(actual)
    path.symlink_to(actual)
    with pytest.raises(ArtifactError, match="symbolic"):
        _find(tmp_path, record, calls)


@pytest.mark.parametrize("extra", ["untrusted", "scopes", "call_id"])
def test_derivation_keys_are_exact(tmp_path: Path, extra: str) -> None:
    record = _fixture(tmp_path)
    record["derived_artifacts"][0][extra] = "submission_B"
    with pytest.raises(ArtifactError, match="fields"):
        validate_derivations(record, tmp_path)


def test_missing_location_duplicate_sources_and_stale_sizes_are_rejected(tmp_path: Path) -> None:
    original = _fixture(tmp_path)
    record = copy.deepcopy(original)
    record["derived_artifacts"][0]["location"] = " "
    with pytest.raises(ArtifactError, match="location"):
        validate_derivations(record, tmp_path)
    record = copy.deepcopy(original)
    record["source_artifacts"].append(record["source_artifacts"][0])
    with pytest.raises(ArtifactError, match="repeats a source"):
        validate_derivations(record, tmp_path)
    record = copy.deepcopy(original)
    record["source_artifacts"][0]["size"] = (tmp_path / _SOURCE).stat().st_size
    record["source_artifacts"][0]["bytes"] = record["source_artifacts"][0]["size"]
    assert validate_derivations(record, tmp_path)
    record["source_artifacts"][0]["bytes"] += 1
    with pytest.raises(ArtifactError, match="size"):
        validate_derivations(record, tmp_path)


def test_structured_derivative_location_is_rejected_under_the_string_contract(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    record["derived_artifacts"][0]["location"] = {"source_page": 21, "derived_page": 1}
    with pytest.raises(ArtifactError, match="location"):
        _find(tmp_path, record, _calls(tmp_path, record))


def test_stream_hash_cannot_be_borrowed_from_other_result_field(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    calls = _calls(tmp_path, record)
    result = calls[_CALL]["result"]
    result["stderr_sha256"] = result["stdout_sha256"]
    result["stdout_sha256"] = "0" * 64
    with pytest.raises(ArtifactError, match="stream hash"):
        _find(tmp_path, record, calls)


def test_receipt_call_id_must_match_trusted_map_key(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    calls = _calls(tmp_path, record)
    calls[_CALL]["call_id"] = "another-call"
    with pytest.raises(ArtifactError, match="call ID"):
        _find(tmp_path, record, calls)


def test_duplicate_json_keys_cannot_hide_declared_source(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    calls = _calls(tmp_path, record)
    text = json.dumps(record).replace('"schema_version": 1', '"schema_version": 0, "schema_version": 1')
    calls[_CALL]["result"]["stdout_sha256"] = _put(tmp_path, _STREAM, text.encode())
    with pytest.raises(ArtifactError, match="repeats an object key"):
        _find(tmp_path, record, calls)


def test_record_size_bound_and_non_json_logs_confer_no_scope(tmp_path: Path) -> None:
    record = _fixture(tmp_path)
    record["checks"][0]["observation"] = "x" * (2 * 1024 * 1024)
    with pytest.raises(ArtifactError, match="2 MiB"):
        validate_derivations(record, tmp_path)
    calls = _calls(tmp_path, record)
    assert _find(tmp_path, record, calls) == set()
    calls[_CALL]["result"]["stdout_sha256"] = _put(tmp_path, _STREAM, b"ordinary process output\n")
    assert _find(tmp_path, record, calls) == set()


def _rendered_pdf_fixture(workspace: Path) -> tuple[dict[str, Any], dict[str, dict[str, Any]], dict[str, Any]]:
    original = "submission/workbook.xlsx"
    derived = ".harness_evidence/scratch/recalculated.pdf"
    source_hash = _put(workspace, original, b"known synthetic workbook bytes")
    pdf_hash = _put(workspace, derived, b"%PDF-1.4 synthetic provenance fixture; no native rendering is performed")
    record = {
        "schema_version": 1,
        "source_artifacts": [{"path": original, "sha256": source_hash}],
        "checks": [
            {
                "action": "Recalculate a scratch copy with a changed input and export its result",
                "observation": {"changed_input": 4, "observed_total": 12, "pages": 1},
            }
        ],
        "derived_artifacts": [{"path": derived, "sha256": pdf_hash, "source_paths": [original], "location": "page 1"}],
    }
    calls = _calls(workspace, record)
    call_id = "000002-rendered-pdf"
    calls[call_id] = {
        "call_id": call_id,
        "tool": "render_pages",
        "arguments": {"path": "scratch/recalculated.pdf", "pages": [1]},
        "failure": None,
        "result": [
            {
                "type": "text",
                "text": json.dumps(
                    {
                        "source": "scratch/recalculated.pdf",
                        "source_sha256": pdf_hash,
                        "pdf_sha256": pdf_hash,
                        "page": 1,
                        "page_count": 1,
                    }
                ),
            }
        ],
    }
    reference = {"path": derived, "sha256": pdf_hash, "call_id": call_id, "method": "visual", "location": "page 1"}
    return record, calls, reference


def test_rendered_scratch_pdf_uses_exact_original_lineage(tmp_path: Path) -> None:
    _, calls, reference = _rendered_pdf_fixture(tmp_path)
    # These are synthetic controller receipts: the test proves provenance
    # routing only, not native rendering or the declared transformation's truth.
    assert _check_ref(reference, tmp_path, calls) == ("visual", {"submission"})


def test_edited_earlier_workbook_copy_does_not_poison_unrelated_pdf_lineage(tmp_path: Path) -> None:
    record, calls, reference = _rendered_pdf_fixture(tmp_path)
    original = tmp_path / "submission/earlier-workbook.xlsx"
    workbook = Workbook()
    workbook.worksheets[0].append([2, 3, "=A1+B1"])
    workbook.save(original)
    scratch = tmp_path / ".harness_evidence/scratch/earlier-workbook.xlsx"
    shutil.copyfile(original, scratch)
    initial_hash = sha256_file(scratch)
    historical = {
        "schema_version": 1,
        "source_artifacts": [{"path": "submission/earlier-workbook.xlsx", "sha256": sha256_file(original)}],
        "checks": [{"action": "Create the initial workbook copy before changing its input", "observation": {"A1": 2}}],
        "derived_artifacts": [
            {
                "path": ".harness_evidence/scratch/earlier-workbook.xlsx",
                "sha256": initial_hash,
                "source_paths": ["submission/earlier-workbook.xlsx"],
                "location": "before editing A1",
            }
        ],
    }
    edited = load_workbook(scratch)
    edited.worksheets[0]["A1"] = 7
    edited.save(scratch)
    assert sha256_file(scratch) != initial_hash
    assert load_workbook(original).worksheets[0]["A1"].value == 2
    with pytest.raises(ArtifactError, match="derived artifact hash"):
        validate_derivations(historical, tmp_path)
    old_id = "000000-initial-workbook-copy"
    old_stream = ".harness_evidence/outputs/initial-workbook-copy.stdout.log"
    old_call = {
        "call_id": old_id,
        "tool": "shell",
        "failure": None,
        "result": {
            "status": "completed",
            "returncode": 0,
            "captured_stdout_path": old_stream,
            "stdout_sha256": _put(tmp_path, old_stream, json.dumps(historical).encode()),
        },
    }
    assert find_derived_scopes(reference["sha256"], tmp_path, {old_id: old_call}, validate_record=_guard) == set()
    assert _check_ref(reference, tmp_path, {old_id: old_call, **calls}) == ("visual", {"submission"})
    # The old declaration is excluded only from the unrelated PDF lookup. The
    # actual source/check/derivative guard still rejects it when it is the target.
    with pytest.raises(ArtifactError, match="derived artifact hash"):
        find_derived_scopes(initial_hash, tmp_path, {old_id: old_call, **calls}, validate_record=_guard)
    assert validate_derivations(record, tmp_path)[0].sha256 == reference["sha256"]


@pytest.mark.parametrize(
    "mutation",
    ["undeclared", "failed_shell", "unrelated_hash", "source_tampered", "derivative_tampered", "scratch_alias"],
)
def test_rendered_scratch_pdf_does_not_infer_unbound_originals(tmp_path: Path, mutation: str) -> None:
    record, calls, reference = _rendered_pdf_fixture(tmp_path)
    if mutation == "undeclared":
        record.pop("derived_artifacts")
        calls.update(_calls(tmp_path, record))
    elif mutation == "failed_shell":
        calls[_CALL]["result"]["returncode"] = 1
    elif mutation == "unrelated_hash":
        other = ".harness_evidence/scratch/other.pdf"
        digest = _put(tmp_path, other, b"another PDF's bytes")
        record["derived_artifacts"][0].update(path=other, sha256=digest)
        calls.update(_calls(tmp_path, record))
    elif mutation == "source_tampered":
        (tmp_path / "submission/workbook.xlsx").write_bytes(b"changed after the original hash")
    elif mutation == "derivative_tampered":
        (tmp_path / reference["path"]).write_bytes(b"PDF changed after the viewed hash")
    else:
        reference["path"] = "scratch/recalculated.pdf"
    if mutation in {"source_tampered", "derivative_tampered", "scratch_alias"}:
        with pytest.raises(ArtifactError, match="hash|missing"):
            _check_ref(reference, tmp_path, calls)
    else:
        assert _check_ref(reference, tmp_path, calls) == ("visual", set())
