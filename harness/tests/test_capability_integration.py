# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from lxml import etree
from openpyxl import Workbook
from pypdf import PdfWriter

from eval_harness.artifacts import file_manifest, manifest_hash, sha256_file, write_json
from eval_harness.capability_environment import parse_environment, prepare_environment
from eval_harness.capability_inspection import InspectionTools
from eval_harness.capability_previews import prepare_shared_previews, stage_shared_previews
from eval_harness.compare_pipeline import _stats
from eval_harness.errors import ArtifactError, ConfigError, HarnessError
from eval_harness.evaluate_pipeline import (
    _inspection_prompt,
    _judge_response_schema,
    _stage_research,
    parse_scalar_score,
)


def test_environment_explicit_and_frozen(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from eval_harness import capability_environment as module

    with pytest.raises(ConfigError, match="network_domains"):
        parse_environment({"profile": "gdpval-v1", "network_domains": ["*"]}, tmp_path)
    with pytest.raises(ConfigError, match="profile"):
        parse_environment({"profile": "invented", "network_domains": []}, tmp_path)
    monkeypatch.setattr(module, "environment_fingerprint", lambda _: "b" * 64)
    with pytest.raises(HarnessError, match="frozen capability environment changed"):
        prepare_environment(
            tmp_path / "work",
            tmp_path,
            {
                "environment": {
                    "profile": "gdpval-v1",
                    "network_domains": [],
                    "fingerprint": "a" * 64,
                }
            },
            "evaluation",
        )


def test_docx_revision_semantics_are_distinct_from_visual_strikethrough(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("The ")
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    deletion = etree.SubElement(paragraph._p, ns + "del", {ns + "id": "1"})
    etree.SubElement(etree.SubElement(deletion, ns + "r"), ns + "delText").text = "old"
    insertion = etree.SubElement(paragraph._p, ns + "ins", {ns + "id": "2"})
    etree.SubElement(etree.SubElement(insertion, ns + "r"), ns + "t").text = "new"
    paragraph.add_run(" story.")
    document.add_paragraph().add_run("Only decoration").font.strike = True
    path = tmp_path / "story.docx"
    document.save(path)
    before = sha256_file(path)
    result = InspectionTools(tmp_path, tmp_path.parent / (tmp_path.name + "-state"), {}).inspect_document("story.docx")
    assert "The old story." in result["original_text"]
    assert "The new story." in result["revised_text"]
    assert result["insertions"] == result["deletions"] == 1
    assert "Only decoration" in result["original_text"]
    assert sha256_file(path) == before


def test_xlsx_tables_are_not_claimed_as_pivot_tables(tmp_path: Path) -> None:
    from openpyxl.worksheet.table import Table

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Name", "Amount"])
    sheet.append(["Operator1", 12])
    sheet.add_table(Table(displayName="Data", ref="A1:B2"))
    sheet["C2"] = "=SUM(B2)"
    workbook.save(tmp_path / "data.xlsx")
    result = InspectionTools(tmp_path, tmp_path.parent / (tmp_path.name + "-state"), {}).inspect_document("data.xlsx")
    assert result["pivot_tables"] == []
    assert result["worksheet_structure"][0]["formula_sample"] == [{"cell": "C2", "formula": "SUM(B2)"}]


def test_inspection_paths_and_page_requests_fail_explicitly(tmp_path: Path) -> None:
    inspector = InspectionTools(tmp_path, tmp_path.parent / (tmp_path.name + "-state"), {})
    outside = tmp_path.parent / (tmp_path.name + "-secret.txt")
    outside.write_text("private")
    (tmp_path / "link.pdf").symlink_to(outside)
    with pytest.raises(ArtifactError, match="escapes"):
        inspector.inspect_document("link.pdf")
    with pytest.raises(ArtifactError, match="one to four"):
        inspector.render_pages("test.pdf", [1, 2, 3, 4, 5])


def test_same_long_pdf_is_available_to_both_judges(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "source"
    source.mkdir()
    pdf = source / "long.pdf"
    writer = PdfWriter()
    for _ in range(30):
        writer.add_blank_page(width=500, height=700)
    writer.write(pdf)

    def copy_pdf(self: InspectionTools, path: Path, target: Path) -> Path:
        target.mkdir(parents=True, exist_ok=True)
        (target / "document.pdf").write_bytes(path.read_bytes())
        write_json(target / "conversion.json", {"page_count": 30})
        return target / "document.pdf"

    monkeypatch.setattr(InspectionTools, "_pdf", copy_pdf)
    cache = tmp_path / "cache"
    prepare_shared_previews([pdf], cache, {})
    images = []
    for name in ("codex", "claude"):
        workspace = tmp_path / name
        (workspace / "submission").mkdir(parents=True)
        (workspace / "submission" / pdf.name).write_bytes(pdf.read_bytes())
        stage_shared_previews(workspace, cache, ["submission"])
        index = json.loads((workspace / ".prepared_previews/index.json").read_text())
        entry = index[sha256_file(pdf)]
        assert entry["page_count"] == 30
        images.append(sha256_file(workspace / ".prepared_previews" / entry["path"]))
    assert images == [sha256_file(pdf)] * 2
    (cache / "index.json").write_text("{}")
    with pytest.raises(ArtifactError, match="changed"):
        prepare_shared_previews([pdf], cache, {})


def test_research_is_staged_without_runtime_identity_and_tampering_is_detected(tmp_path: Path) -> None:
    evidence = tmp_path / "execution/workspace/.harness_evidence"
    (evidence / "research").mkdir(parents=True)
    write_json(evidence / "runtime.json", {"condition": "PRIVATE", "model": "not-for-judge"})
    write_json(evidence / "research/source.json", {"url": "https://example.org/", "kind": "retrieved"})
    entries = file_manifest(evidence)
    expected = manifest_hash(entries)
    write_json(evidence / "manifest.json", {"files": entries})
    destination = tmp_path / "judge/research/submission_A"
    _stage_research(tmp_path / "execution", destination, expected)
    assert [p.name for p in destination.iterdir()] == ["source.json"]
    (evidence / "research/source.json").write_text("changed")
    rewritten = [e for e in file_manifest(evidence) if e["path"] != "manifest.json"]
    write_json(evidence / "manifest.json", {"files": rewritten})
    with pytest.raises(ArtifactError, match="changed"):
        _stage_research(tmp_path / "execution", destination, expected)


@pytest.mark.parametrize("reported", [None, 0, 0.8])
def test_unconfirmed_is_retained_without_becoming_zero_or_a_win(reported: float | None) -> None:
    criteria: dict[str, Any] = {"ai": [{"id": "refresh", "description": "Update behavior"}]}
    response = json.dumps(
        {
            "score": reported,
            "rationale": "Native refresh was not observed",
            "criteria_results": [
                {
                    "id": "refresh",
                    "status": "unconfirmed",
                    "evidence": "Only cached numbers inspected",
                    "reason": "No refresh trial",
                },
            ],
        }
    )
    grade = parse_scalar_score(response, criteria, allow_unconfirmed=True)
    assert grade.valid and grade.score is None
    stats = _stats(
        [
            {
                "task_id": "task",
                "condition_id": "N",
                "repeat": 0,
                "generation_id": "generation",
                "execution_status": "completed",
                "evaluation_status": "completed",
                "score_valid": True,
                "score": grade.score,
                "assessment_status": "unconfirmed",
                "criteria_results": list(grade.criteria_results or ()),
            }
        ],
        "run",
    )
    assert stats["mean_score"] is None
    assert stats["unconfirmed_assessment_count"] == 1
    assert stats["judgment_count"] == stats["generated_sample_count"] == 1
    assert stats["valid_judgment_count"] == 0
    assert _judge_response_schema("scalar", criteria, allow_unconfirmed=True)["properties"]["score"]["type"] == [
        "number",
        "null",
    ]


def test_scratch_authority_is_explicit_in_judge_prompt() -> None:
    from eval_harness.evaluate_pipeline import _pairwise_prompt
    from eval_harness.gdpval import _normalize_row, render_scalar_prompt

    task = _normalize_row({"task_id": "test", "prompt": "Create a workbook."}, Path("source.jsonl"), 1)
    for prompt in (render_scalar_prompt(task), _pairwise_prompt(task)):
        updated = _inspection_prompt(prompt)
        assert "Use scratch copies" in updated
        assert "do not run scripts" not in updated
        assert "score:null" in updated


@pytest.mark.parametrize("reported", ["A", "B", "tie"])
def test_unconfirmed_pairwise_cannot_produce_win_credit(reported: str) -> None:
    from eval_harness.evaluate_pipeline import parse_pairwise_score

    criteria = {"ai": [{"id": "refresh", "description": "Update behavior"}]}
    response = json.dumps(
        {
            "winner": reported,
            "rationale": "Refresh unavailable",
            "criteria_results": [
                {
                    "id": "refresh",
                    "status": "unconfirmed",
                    "evidence": "No refresh observed",
                    "reason": "Missing native path",
                }
            ],
        }
    )
    score = parse_pairwise_score(response, criteria, allow_unconfirmed=True)
    assert score.valid and score.winner == "unjudgeable"
    assert parse_pairwise_score(response, criteria).winner == reported


def test_claude_explicit_mcp_keeps_ambient_customization_disabled(tmp_path: Path) -> None:
    from eval_harness.capability_executor import claude_capability_policy
    from eval_harness.executor import RuntimeSession

    session = RuntimeSession({}, tmp_path, tmp_path)
    policy = claude_capability_policy(session)
    assert policy == {"disableAllHooks": True, "disableBundledSkills": True}
    assert session.environment["CLAUDE_CODE_DISABLE_CLAUDE_MDS"] == "1"
    assert session.environment["ENABLE_CLAUDEAI_MCP_SERVERS"] == "false"
    assert session.environment["DISABLE_PLUGIN_AUTOLOAD"] == "1"


def test_rewritten_preview_manifest_cannot_replace_frozen_evaluation_cache(tmp_path: Path) -> None:
    from eval_harness.config import EvaluationConfig, RuntimeConfig
    from eval_harness.evaluate_pipeline import _prepare_office_previews

    config = EvaluationConfig(
        "scalar",
        RuntimeConfig(
            "codex",
            "explicit-model",
            {"environment": {"profile": "gdpval-v1", "network_domains": [], "office_rendering": {"soffice": "/fake"}}},
        ),
    )
    write_json(tmp_path / "evaluation_manifest.json", {})
    _prepare_office_previews(tmp_path, [], config, tmp_path)
    cache = tmp_path / "inspection_previews"
    write_json(cache / "index.json", {"unexpected-source": {"status": "unavailable"}})
    entries = [e for e in file_manifest(cache) if e["path"] != "manifest.json"]
    write_json(cache / "manifest.json", {"files": entries, "sha256": manifest_hash(entries)})
    with pytest.raises(ArtifactError, match="shared inspection previews changed"):
        _prepare_office_previews(tmp_path, [], config, tmp_path)


def test_verified_shell_does_not_hide_missing_inspection_tools(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import subprocess

    from eval_harness.capability_environment import CapabilityEnvironment
    from eval_harness.capability_executor import probe_environment
    from eval_harness.executor import RuntimeSession

    prepared = CapabilityEnvironment(tmp_path, tmp_path, tmp_path / "config", ("python",), {"specification": {}})
    monkeypatch.setattr(
        subprocess,
        "run",
        lambda *a, **k: subprocess.CompletedProcess(
            a, 0, json.dumps({"verified": True, "tools": ["shell", "fetch", "search"]}), ""
        ),
    )
    with pytest.raises(HarnessError, match="missing required tools"):
        probe_environment(prepared, RuntimeSession({}, tmp_path, tmp_path))


@pytest.mark.parametrize("timed_out", [False, True])
def test_cli_exit_stops_registered_worker_before_snapshot(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, timed_out: bool
) -> None:
    from contextlib import contextmanager
    from types import SimpleNamespace

    from eval_harness import capability_executor as module
    from eval_harness import capability_sandbox
    from eval_harness.capability_environment import CapabilityEnvironment
    from eval_harness.executor import AuthStatus, ExecutionRequest, RuntimeSession

    work, state, config, destination = (tmp_path / name for name in ("work", "state", "config", "destination"))
    for folder in (work, state, config, destination):
        folder.mkdir()
    prepared = CapabilityEnvironment(
        work, state, config / "service.json", ("fake-service",), {"specification": {"network_domains": []}}
    )
    worker = {"alive": False, "stops": 0}

    @contextmanager
    def session(*args: Any):
        yield RuntimeSession({}, work, config)

    class Process:
        returncode = -9 if timed_out else 0
        calls = 0

        def __init__(self, *args: Any, **kwargs: Any) -> None:
            worker["alive"] = True

        def communicate(self, *args: Any, **kwargs: Any) -> tuple[str, str]:
            self.calls += 1
            if timed_out and self.calls == 1:
                raise module.subprocess.TimeoutExpired("fake-cli", 1)
            return '{"type":"turn.completed","usage":{}}\n', ""

    def stop_registered_worker(path: Path) -> None:
        assert path == state
        worker["alive"] = False
        worker["stops"] += 1

    protected = module._protected_manifest

    def snapshot(path: Path, purpose: str) -> list[dict[str, Any]]:
        assert not worker["alive"], "registered worker was still alive while snapshotting protected files"
        return protected(path, purpose)

    def collect(*args: Any, **kwargs: Any) -> None:
        assert not worker["alive"], "registered worker was still alive while collecting deliverables"

    monkeypatch.setattr(module, "prepare_environment", lambda *a: prepared)
    monkeypatch.setattr(module, "probe_environment", lambda *a: {})
    monkeypatch.setattr(module, "_generated_config", lambda *a, **k: "[features]\n")
    monkeypatch.setattr(module.subprocess, "Popen", Process)
    monkeypatch.setattr(module, "_stop_process_group", lambda *a: None)
    monkeypatch.setattr(capability_sandbox, "terminate_active_processes", stop_registered_worker)
    monkeypatch.setattr(module, "_protected_manifest", snapshot)
    monkeypatch.setattr(module, "copy_workspace_deliverables", collect)
    executor = SimpleNamespace(
        _session=session,
        check_auth=lambda *a: AuthStatus(True, True, "synthetic account boundary"),
        _command=lambda *a, **k: ["fake-cli"],
    )
    result = module.execute_with_capabilities(
        executor, ExecutionRequest("synthetic lifecycle probe", destination, "test-model", {}, "application"), "codex"
    )
    assert result.status == ("timeout" if timed_out else "completed")
    assert worker["stops"] > 0 and not worker["alive"]
